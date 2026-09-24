"""Convertit les parties rédigées (Markdown simple) en document Word mis en forme.

Usage :
    python3 md2docx.py sortie.docx partie1.md [partie2.md ...] [--titre-partie "..."] [--sans-page-titre]
"""
import argparse
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
AUTEUR = "Azdine ZARGUI-NAJI"
ENTETE = "Orchestration stratégique RH – Zargui-Naji"
BLEU = RGBColor(0x1F, 0x3A, 0x5F)
NBSP = " "


def set_font(style, size, bold=None, italic=None, color=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    # supprime les polices de thème (sinon les titres s'affichent dans une police sans empattement)
    rfonts = style.element.rPr.rFonts
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(att)) is not None:
            del rfonts.attrib[qn(att)]
    for att in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(att), FONT)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    return run


def typo_fr(text):
    """Espaces insécables de la typographie française."""
    text = text.replace("« ", "«" + NBSP).replace(" »", NBSP + "»")
    return re.sub(r" ([:;?!])", NBSP + r"\1", text)


def add_inline(paragraph, text):
    """Gère **gras** et *italique*."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", typo_fr(text)):
        if not tok:
            continue
        if tok.startswith("**"):
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*"):
            paragraph.add_run(tok[1:-1]).italic = True
        else:
            paragraph.add_run(tok)


def add_table(doc, rows):
    """rows : liste de listes de cellules, la première ligne étant l'en-tête."""
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(3.4), Cm(5.0), Cm(7.6)] if ncol == 3 else [Cm(16 / ncol)] * ncol
    for r, row in enumerate(rows):
        for c, txt in enumerate(row):
            cell = t.cell(r, c)
            cell.width = widths[c]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, txt.strip())
            for run in p.runs:
                run.font.size = Pt(10)
                if r == 0:
                    run.bold = True
            if r == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "DCE3EC")
                cell._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def render_markdown(doc, md_path):
    lines = open(md_path, encoding="utf-8").read().splitlines()
    # démarre au premier titre de section : le bloc de titre Markdown est remplacé par la page de titre
    start = next(i for i, l in enumerate(lines) if l.startswith("## "))
    in_refs = False
    table = []
    for line in lines[start:] + [""]:
        s = line.rstrip()
        if s.startswith("|"):
            cells = s.strip("|").split("|")
            if not all(set(c.strip()) <= set("-:") for c in cells):
                table.append(cells)
            continue
        if table:
            add_table(doc, table)
            table = []
        if not s or s == "---":
            continue
        if s.startswith("### "):
            doc.add_heading(typo_fr(s[4:]), level=2)
        elif s.startswith("## "):
            t = s[3:]
            in_refs = t.lower().startswith("références") or t.lower().startswith("bibliographie")
            if in_refs:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(typo_fr(t), level=1)
        else:
            p = doc.add_paragraph()
            add_inline(p, s)
            if in_refs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Cm(1.25)
                p.paragraph_format.first_line_indent = Cm(-1.25)


def setup_styles(doc):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(2.5))
    normal = doc.styles["Normal"]
    set_font(normal, 12)
    pf = normal.paragraph_format
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in (("Heading 1", 14), ("Heading 2", 12.5), ("Heading 3", 12)):
        st = doc.styles[name]
        set_font(st, size, bold=True, italic=False, color=BLEU)
        st.paragraph_format.space_before = Pt(18 if name == "Heading 1" else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.15
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_title_page(doc, titre_partie):
    def centre(text, size, bold=False, before=0, italic=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic

    centre("CAS Stratégie et organisation – MRHC", 12, before=60)
    centre("Université de Lausanne", 12)
    centre("ORCHESTRATION STRATÉGIQUE RH", 18, bold=True, before=90)
    centre("Audit organisationnel de l'orchestration stratégique RH", 15, before=12)
    centre("La Centrale téléphonique des médecins de garde (CTMG) – Unisanté", 13, italic=True)
    if titre_partie:
        centre(titre_partie, 11, italic=True, before=12)
    centre("Dr Joëlle BÉDAT", 12, before=90)
    centre(AUTEUR, 13, bold=True, before=60)
    centre("Septembre 2026", 12, before=6)
    # la page de titre n'a ni en-tête ni numéro : le corps commence dans une nouvelle section
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.sections[0].header.is_linked_to_previous = False
    doc.sections[0].footer.is_linked_to_previous = False


def add_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.text = ENTETE
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.size = Pt(9)
    hp.runs[0].italic = True
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE").font.size = Pt(10)


def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run("Table des matières")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = BLEU
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    # Word propose de mettre à jour la table des matières à l'ouverture
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)


def build(md_paths, out_path, titre_partie="", page_titre=True):
    doc = Document()
    setup_styles(doc)
    if page_titre:
        add_title_page(doc, titre_partie)
    add_header_footer(doc.sections[-1])
    if page_titre:
        add_toc(doc)
    for md_path in md_paths:
        render_markdown(doc, md_path)
    doc.save(out_path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("sortie")
    ap.add_argument("sources", nargs="+")
    ap.add_argument("--titre-partie", default="")
    ap.add_argument("--sans-page-titre", action="store_true")
    a = ap.parse_args()
    build(a.sources, a.sortie, a.titre_partie, page_titre=not a.sans_page_titre)
