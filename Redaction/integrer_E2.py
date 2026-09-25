"""Intègre l'entretien E2 anonymisé dans l'annexe C2 d'une version Word du devoir.

Usage : python3 integrer_E2.py entree.docx entretien_E2.docx sortie.docx
"""
import re
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NBSP = " "
LABELS = {"Azdine": "Auteur", "Corinne": "E2", "[Intervenant 2]": "Auteur"}
RETRAIT = "[Données personnelles retirées pour préserver l'anonymat.]"


def anonymise(text):
    text = text.replace("J'ai une femme, j'ai 59 ans et j'ai une française.", RETRAIT)
    text = re.sub(r"\bCorinne\b", "[E2]", text)
    return re.sub(r"\bAzdine\b", "[Auteur]", text)


def lire_tours(src):
    """Retourne (tours de l'autorisation, tours de l'entretien) : listes de (locuteur, [paragraphes])."""
    paras = [p.text.strip() for p in Document(src).paragraphs]
    debut = next(i for i, t in enumerate(paras) if t.startswith("Restitution"))
    autorisation, entretien = [], []
    for t in paras[:debut]:
        m = re.match(r"^(Azdine|Corinne)\s*:\s*(.+)$", t.replace(NBSP, " "))
        if m:
            autorisation.append((LABELS[m.group(1)], [anonymise(m.group(2))]))
    courant = None
    for t in paras[debut + 1:]:
        brut = t.replace(NBSP, " ").strip()
        if not brut:
            continue
        m = re.match(r"^(Azdine|Corinne)\s*:\s*$", brut)
        if m or brut == "[Intervenant 2]":
            courant = (LABELS[m.group(1) if m else brut], [])
            entretien.append(courant)
        elif courant is not None:
            courant[1].append(anonymise(brut))
    return autorisation, [t for t in entretien if t[1]]


def inserer_avant(ref, texte="", label=None, italique=False, gras=False):
    p = ref.insert_paragraph_before("", style=ref.style if ref.style.name == "Normal" else "Normal")
    if label:
        p.add_run(label + NBSP + ":").bold = True
        p.add_run(" " + texte)
    else:
        r = p.add_run(texte)
        r.italic = italique
        r.bold = gras
    return p


def main(src, entretien_src, dst):
    doc = Document(src)
    paras = doc.paragraphs
    i_c2 = next(i for i, p in enumerate(paras) if p.text.startswith("C2. Entretien"))
    notice = paras[i_c2 + 1]
    suivant = paras[i_c2 + 2]  # titre de l'annexe D
    assert suivant.text.startswith("Annexe D"), suivant.text

    # notice de l'annexe C2, sur le modèle de C1
    for r in notice.runs[1:]:
        r.text = ""
    notice.runs[0].text = (
        "Date" + NBSP + ": 28.11.2024. Cadre" + NBSP + ": entretien téléphonique réalisé selon la technique des "
        "incidents critiques pour le module 002 (Comportements organisationnels), consentement recueilli en début "
        "d'entretien. Restitution" + NBSP + ": retranscription automatique (Turboscribe), corrigée et complétée par "
        "réécoute. Autorisation de réutilisation pour le présent audit" + NBSP + ": accord oral donné par téléphone "
        "le 17.09.2026 (échange retranscrit ci-dessous). Les prénoms ont été remplacés par «" + NBSP + "Auteur" + NBSP
        + "» et «" + NBSP + "E2" + NBSP + "», et les données personnelles d'identification ont été retirées."
    )
    notice.runs[0].italic = True

    autorisation, entretien = lire_tours(entretien_src)
    inserer_avant(suivant, "Échange du 17.09.2026 – autorisation de réutilisation", gras=True)
    for loc, textes in autorisation:
        inserer_avant(suivant, textes[0], label=loc)
    inserer_avant(suivant, "Entretien du 28.11.2024", gras=True)
    for loc, textes in entretien:
        inserer_avant(suivant, textes[0], label=loc)
        for t in textes[1:]:
            inserer_avant(suivant, t)
    inserer_avant(suivant, "")

    # Word proposera de mettre à jour le sommaire (numéros de page des annexes) à l'ouverture
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)
    doc.save(dst)
    print(len(autorisation), "tours (autorisation),", len(entretien), "tours (entretien)")


if __name__ == "__main__":
    main(*sys.argv[1:4])
