"""Applique les corrections acceptées par l'auteur (sans mode révision) et retire quatre références.

Usage : python3 appliquer_corrections_finales.py entree_fusionnee.docx sortie.docx
(entree_fusionnee.docx : version finale de l'auteur passée par merge_runs.py)
"""
import sys

from docx import Document

NBSP = " "

# (début du paragraphe, texte à remplacer, remplacement) – corrections acceptées
REMPLACEMENTS = [
    ("Le cadre d", "grille d’évaluation du module", "grille d’audit du module"),
    ("L’analyse est basée", "L’analyse est basée sur la perspective de la fonction de chef de salle à la CTMG.",
     "L’analyse adopte la perspective de l’auteur, chef de salle à la CTMG."),
    ("L’analyse ", "Ne disposant pas de l’accès à ces données, ils ne peuvent pas être mesurés, et nous les présentons comme tels.",
     "Faute d’accès à ces données, nous ne pouvons pas les mesurer, et nous les présentons comme tels."),
    ("Le matériau mobilisé", "Comportements Organisationnels", "Comportements organisationnels"),
    ("Le matériau mobilisé", "Ce choix découle toujours de la contrainte", "Ce choix découle de la contrainte"),
    ("Cette proximité expose",
     "afin de rester dans le cadre d’un audit adressé à un conseil ayant vocation à décrire des mécanismes de l’organisation",
     "car un audit adressé à un conseil a vocation à décrire des mécanismes de l’organisation"),
    ("Depuis le 1er juillet 2024", "la CTMG intègre comme secteur à part entière du Département",
     "la CTMG est intégrée, comme secteur à part entière, au Département"),
    ("L'enjeu principal", "24/7", "24 heures sur 24 et 7 jours sur 7"),
    ("L'enjeu principal", "repose notamment par la qualité", "repose notamment sur la qualité"),
    ("L'enjeu principal", "de santé publique de régulateurs", "de santé publique des régulateurs"),
    ("L'enjeu principal", "outils de supports", "outils de support"),
    ("L'enjeu principal",
     "afin d’analyser et orienter les appelants caractérisées par une grande variété de situations et problématiques",
     "afin d’analyser et d’orienter des appels caractérisés par une grande variété de situations et de problématiques"),
    ("Une gouvernance à distance", " qui peut être assimilé à la", ", ce qui renvoie à la"),
    ("Au sein de la CTMG, une ligne",
     "Une position rendant ces derniers transparents dans la communication ascendante et silencieux dans les prises de décisions.",
     "Cette position les rend transparents dans la communication ascendante et sans voix dans les prises de décision."),
    ("À la CTMG, les canaux", "ainsi que d’observation et gestion managériale limitée",
     "ainsi que d’observations managériales limitées"),
    ("L'expertise des régulateurs", "Ces faits exposent donc un risque", "Ces faits exposent donc l’organisation à un risque"),
    ("Une centrale en sous-effectif", "Ces conditions exposent un risque", "Ces conditions exposent à un risque"),
    ("Ce travail repose", "sur les observations en qualité de chef de salle", "sur nos observations en qualité de chef de salle"),
    # retrait des références non exploitées (Kunz, Lima & Galleli, Martin et al., Gooderham et al.)
    ("Jungmann (2006)", "des configurations variées (Kunz, 2010), et", "des configurations variées, et"),
    ("La direction des ressources humaines",
     "Lima et Galleli (2021) ainsi que Martin et al. (2016) montrent que l'articulation entre gouvernance et GRH "
     "stratégique peut prendre des formes très différentes selon la place que la gouvernance accorde au capital humain. ",
     ""),
    ("Reichel et Lazarova",
     "Gooderham et al. (2015) montrent que cette délégation dépend autant de facteurs nationaux que de choix propres à "
     "l'organisation. ", ""),
]

REFERENCES_A_RETIRER = ("Gooderham, P. N.", "Kunz, P. V.", "Lima, L., & Galleli", "Martin, G., Farndale")


def norm(s):
    return s.replace(NBSP, " ").replace(" ", " ")


def remplacer(paragraph, old, new):
    for run in paragraph.runs:
        t = run.text
        i = norm(t).find(norm(old))
        if i >= 0:
            run.text = t[:i] + new + t[i + len(old):]
            return True
    return False


def main(src, dst):
    doc = Document(src)
    ko = []
    for start, old, new in REMPLACEMENTS:
        cibles = [p for p in doc.paragraphs
                  if norm(p.text).lstrip().startswith(norm(start)) and norm(old) in norm(p.text)]
        if not (cibles and remplacer(cibles[0], old, new)):
            ko.append(old[:70])
    retirees = 0
    for p in list(doc.paragraphs):
        if p.text.startswith(REFERENCES_A_RETIRER):
            p._p.getparent().remove(p._p)
            retirees += 1
    doc.save(dst)
    print(len(REMPLACEMENTS) - len(ko), "remplacements appliqués ;", retirees, "références retirées")
    for k in ko:
        print("NON APPLIQUÉ :", k)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
